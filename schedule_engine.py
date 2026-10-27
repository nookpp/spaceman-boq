import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Dict, Any, List, Optional
from openai import OpenAI

DEFAULT_BASE_URL = "https://opencode.ai/zen/go/v1"
DEFAULT_MODEL = "deepseek-v4.1-flash"

SCHEDULE_SYSTEM_PROMPT = """คุณคือผู้เชี่ยวชาญการวางแผนโครงการออกแบบและตกแต่งภายในของ SPACEMAN STUDIO
หน้าที่ของคุณคือ:
1. วิเคราะห์ข้อความหรือความต้องการไทม์ไลน์งานที่ผู้ใช้พิมพ์บอก
2. แปลงออกมาเป็นโครงสร้าง JSON ของ Project Schedule (Gantt Chart) ตามมาตรฐาน SPACEMAN STUDIO
3. จัดแบ่งสเตจงานมาตรฐาน เช่น:
   - A. Schematic Design
   - B. Preliminary design
   - C. Design develop
   - D. Final drawing 1
   - E. Final drawing 2
   - F. BOQ stage
   - G. Construction stage
4. กำหนดเดือนและสัปดาห์ (แต่ละเดือนมี 4 สัปดาห์: Week 1, 2, 3, 4)
5. ระบุระยะเวลา (Duration in weeks), สัปดาห์เริ่มต้น-สิ้นสุด, สัปดาห์ที่มีนัดประชุมลูกค้า (สัญลักษณ์ M), และสัปดาห์เปิดร้าน (OPENING DATE)

ต้องส่งผลลัพธ์เป็น JSON ล้วนๆ ในรูปแบบ:
{
  "project_name": "ชื่อโครงการ เช่น Brewave Aree",
  "date": "วันที่ออกไทม์ไลน์ เช่น 09 May 2025",
  "year": 2025,
  "months": ["May", "June", "July", "Aug", "Sep", "Oct", "Nov"],
  "opening_date": {"month": "Oct", "week": 2},
  "phases": [
    {
      "code": "A",
      "title": "Schematic Design",
      "color": "#FFF59D",
      "tasks": [
        {
          "description": "Lay-out",
          "duration": 1,
          "start_month": "May",
          "start_week": 2,
          "end_month": "May",
          "end_week": 2,
          "meetings": [{"month": "May", "week": 2}]
        }
      ]
    }
  ]
}
อย่าพิมพ์ข้อความเกริ่นหรือลงท้ายใดๆ ให้ส่งเฉพาะ JSON เท่านั้น
"""

def get_sample_brewave_schedule() -> Dict[str, Any]:
    """Returns exact pre-calibrated schedule matching Brewave Aree sample image."""
    return {
        "project_name": "Brewave Aree",
        "date": "09 May 2025",
        "year": 2025,
        "months": ["May", "June", "July", "Aug", "Sep", "Oct", "Nov"],
        "opening_date": {"month": "Oct", "week": 2},
        "phases": [
            {
                "code": "A",
                "title": "Schematic Design",
                "color": "#FFF59D",
                "tasks": [
                    {
                        "description": "Lay-out",
                        "duration": 1,
                        "start_month": "May",
                        "start_week": 2,
                        "end_month": "May",
                        "end_week": 2,
                        "meetings": [{"month": "May", "week": 2}]
                    }
                ]
            },
            {
                "code": "B",
                "title": "Preliminary design",
                "color": "#C8E6C9",
                "tasks": [
                    {"description": "Rendering", "duration": 1, "start_month": "May", "start_week": 3, "end_month": "May", "end_week": 3, "meetings": [{"month": "May", "week": 3}]},
                    {"description": "Revision stage", "duration": 1, "start_month": "May", "start_week": 4, "end_month": "May", "end_week": 4, "meetings": []},
                    {"description": "Final rendering", "duration": 1, "start_month": "June", "start_week": 1, "end_month": "June", "end_week": 1, "meetings": [{"month": "June", "week": 1}]},
                    {"description": "Materials presentation", "duration": 2, "start_month": "June", "start_week": 1, "end_month": "June", "end_week": 2, "meetings": [{"month": "June", "week": 1}]}
                ]
            },
            {
                "code": "C",
                "title": "Design develop",
                "color": "#BBDEFB",
                "tasks": [
                    {"description": "MEP Drawing for tender", "duration": 2, "start_month": "June", "start_week": 1, "end_month": "June", "end_week": 2, "meetings": []},
                    {"description": "AC Drawing for tender", "duration": 2, "start_month": "June", "start_week": 1, "end_month": "June", "end_week": 2, "meetings": []}
                ]
            },
            {
                "code": "D",
                "title": "Final drawing 1",
                "color": "#00E5FF",
                "tasks": [
                    {"description": "Interior joinery drawing for tender", "duration": 4, "start_month": "June", "start_week": 1, "end_month": "June", "end_week": 4, "meetings": []}
                ]
            },
            {
                "code": "E",
                "title": "Final drawing 2",
                "color": "#90CAF9",
                "tasks": [
                    {"description": "Interior detail drawing & equipment specs", "duration": 4, "start_month": "June", "start_week": 4, "end_month": "July", "end_week": 3, "meetings": []}
                ]
            },
            {
                "code": "F",
                "title": "BOQ stage",
                "color": "#FFCC80",
                "tasks": [
                    {"description": "MEP/AC BOQ", "duration": 1, "start_month": "June", "start_week": 2, "end_month": "June", "end_week": 2, "meetings": [{"month": "June", "week": 2}]},
                    {"description": "Interior construction BOQ", "duration": 3, "start_month": "July", "start_week": 1, "end_month": "July", "end_week": 3, "meetings": [{"month": "July", "week": 3}]},
                    {"description": "Loose furniture ordering", "duration": 10, "start_month": "July", "start_week": 3, "end_month": "Sep", "end_week": 4, "meetings": [{"month": "July", "week": 3}, {"month": "Aug", "week": 3}, {"month": "Sep", "week": 3}]}
                ]
            },
            {
                "code": "G",
                "title": "Construction stage",
                "color": "#F48FB1",
                "tasks": [
                    {"description": "MEP/AC Construction", "duration": 6, "start_month": "June", "start_week": 3, "end_month": "July", "end_week": 4, "meetings": [{"month": "June", "week": 3}, {"month": "July", "week": 3}]},
                    {"description": "Interior construction", "duration": 10, "start_month": "July", "start_week": 3, "end_month": "Sep", "end_week": 4, "meetings": [{"month": "July", "week": 3}, {"month": "Aug", "week": 2}, {"month": "Sep", "week": 2}, {"month": "Sep", "week": 4}]},
                    {"description": "Interior defected Maintenance", "duration": 4, "start_month": "Oct", "start_week": 3, "end_month": "Nov", "end_week": 2, "meetings": [{"month": "Oct", "week": 3}, {"month": "Nov", "week": 2}]}
                ]
            }
        ]
    }

def parse_schedule_prompt(
    api_key: str,
    user_prompt: str,
    base_url: str = DEFAULT_BASE_URL,
    model_name: str = DEFAULT_MODEL
) -> Dict[str, Any]:
    """Uses SPACEMAN AI Engine to extract structured schedule JSON from natural text."""
    if not api_key:
        raise ValueError("กรุณาระบุ AI Engine API Key")

    client = OpenAI(base_url=base_url, api_key=api_key, default_headers={"x-opencode-session": "ses_spaceman_schedule"})
    
    response = client.chat.completions.create(
        model=model_name,
        messages=[
            {"role": "system", "content": SCHEDULE_SYSTEM_PROMPT},
            {"role": "user", "content": f"สร้าง Project Schedule ไทม์ไลน์งานจากข้อความต่อไปนี้:\n\n{user_prompt}"}
        ],
        temperature=0.2
    )

    raw_content = response.choices[0].message.content or "{}"
    cleaned = raw_content.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.split("\n", 1)[-1]
    if cleaned.endswith("```"):
        cleaned = cleaned.rsplit("```", 1)[0]
    cleaned = cleaned.strip()

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        import re
        match = re.search(r'\{.*\}', cleaned, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError(f"AI ตอบกลับมาในรูปแบบที่ไม่ใช่ JSON ถูกต้อง: {raw_content[:200]}")

def generate_schedule_excel(data: Dict[str, Any]) -> io.BytesIO:
    """Creates a beautifully styled Excel Gantt chart matching SPACEMAN STUDIO format."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Schedule"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    thin_border_side = Side(border_style="thin", color="000000")
    double_border_side = Side(border_style="double", color="000000")
    grid_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    font_title = Font(name="Calibri", size=14, bold=True)
    font_subtitle = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_m = Font(name="Calibri", size=10, bold=True, color="000000")
    font_opening = Font(name="Calibri", size=9, bold=True, color="FFFFFF")

    # Header Titles
    project_name = data.get("project_name", "Project Name")
    ws["B2"] = f"Project Name : {project_name}"
    ws["B2"].font = font_title
    
    ws["B4"] = "Project Schedule"
    ws["B4"].font = font_subtitle
    
    date_str = data.get("date", "09 May 2025")
    ws["B5"] = f"Date : {date_str}"
    ws["B5"].font = font_bold

    # Spaceman Studio Brand in top right
    months = data.get("months", ["May", "June", "July", "Aug", "Sep", "Oct", "Nov"])
    total_weeks = len(months) * 4
    brand_col = 4 + total_weeks
    brand_cell = ws.cell(2, brand_col - 2)
    brand_cell.value = "SPACEMAN STUDIO"
    brand_cell.font = Font(name="Arial", size=14, bold=True, color="1E293B")
    brand_cell.alignment = Alignment(horizontal="right")

    # Timeline Column Headers (Row 7, 8, 9)
    # Col B: No. (2)
    # Col C: Description (3)
    # Col D: Duration (week) (4)
    # Col E onwards: Weeks (5 to 5 + total_weeks - 1)
    ws.merge_cells("B7:B9")
    ws["B7"] = "No."
    ws.merge_cells("C7:C9")
    ws["C7"] = "Description"
    ws.merge_cells("D7:D9")
    ws["D7"] = "Deration\n(week)"
    
    for c in ["B7", "C7", "D7"]:
        ws[c].font = font_bold
        ws[c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[c].fill = PatternFill(start_color="F2F2F2", fill_type="solid")

    # Year Row (Row 7)
    start_timeline_col = 5
    end_timeline_col = start_timeline_col + total_weeks - 1
    
    year_cell = ws.cell(7, start_timeline_col)
    year_cell.value = str(data.get("year", 2025))
    year_cell.font = font_bold
    year_cell.alignment = Alignment(horizontal="center", vertical="center")
    year_cell.fill = PatternFill(start_color="E0E0E0", fill_type="solid")
    ws.merge_cells(start_row=7, start_column=start_timeline_col, end_row=7, end_column=end_timeline_col)

    # Month Row (Row 8) & Week Row (Row 9)
    col_idx = start_timeline_col
    week_col_map = {} # (month_name, week_num) -> col_index

    for m_idx, m_name in enumerate(months):
        m_start_col = col_idx
        m_end_col = col_idx + 3
        ws.merge_cells(start_row=8, start_column=m_start_col, end_row=8, end_column=m_end_col)
        m_cell = ws.cell(8, m_start_col)
        m_cell.value = m_name
        m_cell.font = font_bold
        m_cell.alignment = Alignment(horizontal="center", vertical="center")
        m_cell.fill = PatternFill(start_color="ECEFF1", fill_type="solid")
        
        for w in range(1, 5):
            w_cell = ws.cell(9, col_idx)
            w_cell.value = w
            w_cell.font = font_regular
            w_cell.alignment = Alignment(horizontal="center", vertical="center")
            w_cell.fill = PatternFill(start_color="F5F5F5", fill_type="solid")
            week_col_map[(m_name.lower(), w)] = col_idx
            col_idx += 1

    # Border header rows 7-9
    for r in range(7, 10):
        for c in range(2, end_timeline_col + 1):
            ws.cell(r, c).border = grid_border

    # Data Rows (Row 10 onwards)
    current_row = 10
    phases = data.get("phases", [])
    opening_info = data.get("opening_date", {})
    opening_col = None
    if opening_info:
        o_month = str(opening_info.get("month", "")).lower()
        o_week = int(opening_info.get("week", 2))
        opening_col = week_col_map.get((o_month, o_week))

    for phase in phases:
        # Phase Header Row (e.g. A. Schematic Design)
        p_code = phase.get("code", "")
        p_title = phase.get("title", "")
        p_color = phase.get("color", "#E0E0E0").replace("#", "")
        
        ws.cell(current_row, 2).value = f"{p_code}."
        ws.cell(current_row, 2).font = font_bold
        ws.cell(current_row, 2).alignment = Alignment(horizontal="center", vertical="center")
        
        ws.cell(current_row, 3).value = p_title
        ws.cell(current_row, 3).font = font_bold
        ws.cell(current_row, 3).alignment = Alignment(vertical="center")
        
        for c in range(2, end_timeline_col + 1):
            cell = ws.cell(current_row, c)
            cell.border = grid_border
            cell.fill = PatternFill(start_color="F8FAFC", fill_type="solid")

        current_row += 1

        # Tasks
        for task in phase.get("tasks", []):
            desc = task.get("description", "")
            duration = task.get("duration", 1)
            s_month = str(task.get("start_month", "")).lower()
            s_week = int(task.get("start_week", 1))
            e_month = str(task.get("end_month", "")).lower()
            e_week = int(task.get("end_week", 1))
            meetings = [(str(m.get("month", "")).lower(), int(m.get("week", 1))) for m in task.get("meetings", [])]

            ws.cell(current_row, 2).value = "-"
            ws.cell(current_row, 2).alignment = Alignment(horizontal="center", vertical="center")
            
            ws.cell(current_row, 3).value = desc
            ws.cell(current_row, 3).font = font_regular
            ws.cell(current_row, 3).alignment = Alignment(vertical="center")
            
            ws.cell(current_row, 4).value = duration
            ws.cell(current_row, 4).font = font_regular
            ws.cell(current_row, 4).alignment = Alignment(horizontal="center", vertical="center")

            # Determine start col and end col
            start_c = week_col_map.get((s_month, s_week))
            end_c = week_col_map.get((e_month, e_week))
            
            if start_c is None:
                start_c = start_timeline_col
            if end_c is None:
                end_c = min(start_c + duration - 1, end_timeline_col)

            # Draw timeline cells for this task
            for c in range(start_timeline_col, end_timeline_col + 1):
                cell = ws.cell(current_row, c)
                cell.border = grid_border
                
                # Check if this cell is within task duration bar
                if start_c <= c <= end_c:
                    cell.fill = PatternFill(start_color=p_color, fill_type="solid")
                    
                    # Check if there is a meeting 'M'
                    for m_month, m_week in meetings:
                        if week_col_map.get((m_month, m_week)) == c:
                            cell.value = "M"
                            cell.font = font_m
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                            break
                else:
                    # Empty cell
                    cell.fill = PatternFill(fill_type=None)

            # Also border cols 2, 3, 4
            ws.cell(current_row, 2).border = grid_border
            ws.cell(current_row, 3).border = grid_border
            ws.cell(current_row, 4).border = grid_border

            current_row += 1

    # Draw OPENING DATE vertical column if specified
    if opening_col:
        for r in range(10, current_row):
            cell = ws.cell(r, opening_col)
            cell.fill = PatternFill(start_color="37474F", fill_type="solid")
            cell.font = font_opening
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Label OPENING DATE in the middle row of that column
        mid_row = 10 + (current_row - 10) // 2
        ws.cell(mid_row, opening_col).value = "OPENING DATE"
        ws.cell(mid_row, opening_col).alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)

    # Footer Remark
    current_row += 1
    ws.cell(current_row, 2).value = "Remark** M= สัปดาห์ที่มีการนัดประชุม"
    ws.cell(current_row, 2).font = Font(name="Calibri", size=10, italic=True)

    # Column Widths
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 12
    for c in range(start_timeline_col, end_timeline_col + 2):
        col_letter = openpyxl.utils.get_column_letter(c)
        ws.column_dimensions[col_letter].width = 4.2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def render_schedule_html(data: Dict[str, Any]) -> str:
    """Generates a responsive HTML table Gantt chart matching SPACEMAN STUDIO aesthetic."""
    project_name = data.get("project_name", "Project Name")
    date_str = data.get("date", "09 May 2025")
    year = data.get("year", 2025)
    months = data.get("months", ["May", "June", "July", "Aug", "Sep", "Oct", "Nov"])
    opening_info = data.get("opening_date", {})
    o_month = str(opening_info.get("month", "")).lower() if opening_info else ""
    o_week = int(opening_info.get("week", 2)) if opening_info else 0

    # Build header HTML
    html = f"""
    <div style="font-family: Arial, sans-serif; background: #ffffff; padding: 20px; border-radius: 8px; border: 1px solid #e2e8f0; overflow-x: auto;">
        <div style="display: flex; justify-content: space-between; align-items: flex-start; margin-bottom: 15px;">
            <div>
                <h3 style="margin: 0 0 5px 0; color: #0f172a; font-size: 1.3rem;">Project Name : {project_name}</h3>
                <div style="font-weight: bold; color: #334155; font-size: 1rem;">Project Schedule</div>
                <div style="font-size: 0.9rem; color: #64748b; margin-top: 3px;">Date : {date_str}</div>
            </div>
            <div style="font-size: 1.4rem; font-weight: 800; color: #0f172a; letter-spacing: 1px; display: flex; align-items: center; gap: 8px;">
                <span style="font-size: 1.6rem;">🪐</span> SPACEMAN STUDIO
            </div>
        </div>

        <table style="border-collapse: collapse; width: 100%; min-width: 900px; font-size: 12px; text-align: center;">
            <thead>
                <tr style="background: #f1f5f9; color: #1e293b;">
                    <th rowspan="3" style="border: 1px solid #cbd5e1; padding: 6px; width: 45px;">No.</th>
                    <th rowspan="3" style="border: 1px solid #cbd5e1; padding: 6px; text-align: left; min-width: 260px;">Description</th>
                    <th rowspan="3" style="border: 1px solid #cbd5e1; padding: 6px; width: 70px;">Deration<br>(week)</th>
                    <th colspan="{len(months) * 4}" style="border: 1px solid #cbd5e1; padding: 4px; background: #e2e8f0; font-weight: bold;">{year}</th>
                </tr>
                <tr style="background: #f8fafc; color: #334155; font-weight: bold;">
    """
    for m in months:
        html += f'<th colspan="4" style="border: 1px solid #cbd5e1; padding: 4px;">{m}</th>'
    html += '</tr><tr style="background: #f1f5f9; color: #64748b; font-size: 11px;">'

    # Week numbers
    all_week_keys = []
    for m in months:
        for w in range(1, 5):
            all_week_keys.append((m.lower(), w))
            html += f'<th style="border: 1px solid #cbd5e1; padding: 3px; width: 24px;">{w}</th>'
    html += "</tr></thead><tbody>"

    # Phases and Tasks
    for phase in data.get("phases", []):
        p_code = phase.get("code", "")
        p_title = phase.get("title", "")
        p_color = phase.get("color", "#c8e6c9")

        # Phase Header
        html += f"""
        <tr style="background: #f8fafc; font-weight: bold; color: #0f172a;">
            <td style="border: 1px solid #cbd5e1; padding: 5px;">{p_code}.</td>
            <td colspan="{2 + len(all_week_keys)}" style="border: 1px solid #cbd5e1; padding: 5px; text-align: left;">{p_title}</td>
        </tr>
        """

        for task in phase.get("tasks", []):
            desc = task.get("description", "")
            duration = task.get("duration", 1)
            s_m = str(task.get("start_month", "")).lower()
            s_w = int(task.get("start_week", 1))
            e_m = str(task.get("end_month", "")).lower()
            e_w = int(task.get("end_week", 1))
            meetings = [(str(m.get("month", "")).lower(), int(m.get("week", 1))) for m in task.get("meetings", [])]

            s_idx = -1
            e_idx = -1
            for idx, wk in enumerate(all_week_keys):
                if wk == (s_m, s_w) and s_idx == -1:
                    s_idx = idx
                if wk == (e_m, e_w):
                    e_idx = idx
            if s_idx == -1:
                s_idx = 0
            if e_idx == -1:
                e_idx = min(s_idx + duration - 1, len(all_week_keys) - 1)

            html += f"""
            <tr style="height: 24px;">
                <td style="border: 1px solid #cbd5e1; color: #94a3b8;">-</td>
                <td style="border: 1px solid #cbd5e1; text-align: left; padding: 4px 8px; color: #1e293b;">{desc}</td>
                <td style="border: 1px solid #cbd5e1; color: #475569;">{duration}</td>
            """

            for idx, wk in enumerate(all_week_keys):
                # Is Opening date column?
                is_opening = (wk == (o_month, o_week))
                
                # Is in task duration?
                in_task = (s_idx <= idx <= e_idx)
                
                # Is meeting?
                is_meeting = (wk in meetings and in_task)
                m_text = "M" if is_meeting else ""

                if is_opening:
                    html += f'<td style="border: 1px solid #cbd5e1; background: #334155; color: #ffffff; font-weight: bold; font-size: 10px;">{m_text}</td>'
                elif in_task:
                    html += f'<td style="border: 1px solid #cbd5e1; background: {p_color}; color: #000000; font-weight: bold;">{m_text}</td>'
                else:
                    html += '<td style="border: 1px solid #cbd5e1;"></td>'

            html += "</tr>"

    html += f"""
            </tbody>
        </table>
        <div style="margin-top: 10px; font-size: 0.85rem; color: #475569; font-style: italic;">
            Remark** M = สัปดาห์ที่มีการนัดประชุม | แถบสีเข้ม = OPENING DATE (เป้าหมายเปิดร้าน)
        </div>
    </div>
    """
    return html

